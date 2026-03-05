from googleapiclient.discovery import build
from vertexai import rag
import vertexai
from vertexai.preview.generative_models import GenerativeModel, Tool
from google.cloud import aiplatform
from rag_agent.config import DEFAULT_EMBEDDING_MODEL, MODEL, DEFAULT_CHUNK_OVERLAP, DEFAULT_CHUNK_SIZE, DEFAULT_EMBEDDING_REQUESTS_PER_MIN

drive_service = build('drive', 'v3')
# folder_query = f"name = 'ACCOR' and mimeType = 'application/vnd.google-apps.folder' and trashed = false"
# folder_results = drive_service.files().list(q=folder_query, fields="files(id, name)").execute()
# folders = folder_results.get('files', [])

# # Test simple pour lister TOUT ce que le Service Account voit
# results = drive_service.files().list(pageSize=10).execute()
# files = results.get('files', [])
# print(f"Le compte de service voit {len(files)} fichiers/dossiers.")

vertexai.init(project="knowledge-base-agent-485813", location="europe-west4")

embedding_model_config = rag.RagEmbeddingModelConfig(
            vertex_prediction_endpoint=rag.VertexPredictionEndpoint(
                publisher_model=DEFAULT_EMBEDDING_MODEL
            )
        )

rag_corpus = rag.create_corpus(
            display_name="test",
            backend_config=rag.RagVectorDbConfig(
                rag_embedding_model_config=embedding_model_config
            ),
        )

chunking_config = rag.ChunkingConfig(
    chunk_size=512,        # Nombre de tokens par morceau
    chunk_overlap=50       # Chevauchement entre les morceaux
)

transformation_config = rag.TransformationConfig(
    chunking_config=chunking_config
)

rag.import_files(
    corpus_name=rag_corpus.name,
    paths=["https://drive.google.com/drive/u/2/folders/1DujBsU1kC-Oye4M1tDTBsITpYeBHxsJE"],
    transformation_config=transformation_config
)

# response = rag.retrieval_query(
#     rag_resources=[
#         rag.RagResource(
#             rag_corpus=rag_corpus.name,
#         )
#     ],
#     text="Quelle est la roadmap de Backmarket ?",
# )

# print(response)

rag_store = rag.VertexRagStore(
    rag_resources=[
        rag.RagResource(
            rag_corpus.name
            )
        ],
    rag_retrieval_config= rag.RagRetrievalConfig(
                top_k=10,
                filter=rag.utils.resources.Filter(vector_distance_threshold=0.5),
            )
)

rag_retrieval_tool = Tool.from_retrieval(
    retrieval=rag.Retrieval(source=rag_store)
)

rag_model = GenerativeModel(
    model_name="gemini-2.5-flash", tools=[rag_retrieval_tool]
)

response = rag_model.generate_content("Quelle est la roadmap de Backmarket ? Donne les liens des fichiers où tu as trouvé les infos.")
print(response.text)

if response.candidates[0].grounding_metadata:
            metadata = response.candidates[0].grounding_metadata
            # Extraction des URIs uniques pour ne pas polluer l'agent
            sources_set = set()
            for chunk in metadata.grounding_chunks:
                if hasattr(chunk, 'retrieved_context'):
                    sources_set.add(chunk.retrieved_context.uri)
        
sources_info = []
if response.candidates[0].grounding_metadata:
    metadata = response.candidates[0].grounding_metadata
    seen_uris = set()
    
    for chunk in metadata.grounding_chunks:
        # Récupération sécurisée du contexte de retrieval
        ctx = getattr(chunk, 'retrieved_context', None)
        uri = getattr(ctx, 'uri', None) if ctx else None
        title = getattr(ctx, 'title', None) if ctx else "Document sans titre"
        
        if uri and uri not in seen_uris:
            seen_uris.add(uri)
            sources_info.append({
                "title": title,
                "uri": uri
            })
            
response = rag_model.generate_content("Donne moi tous les fichiers qui parlent de la roadmap")
print(response.text)

response = rag_model.generate_content("Donne moi le fichier le plus récent et sa date de dernière mise a jour")
print(response.text)

response = rag_model.generate_content("Donne moi le lien du fichier le plus récent et sa date de dernière mise a jour")
print(response.text)




folder_query = f"name = 'BACKMARKET' and mimeType = 'application/vnd.google-apps.folder' and trashed = false"
folder_results = drive_service.files().list(q=folder_query, fields="files(id, name)").execute()
folders = folder_results.get('files', [])



from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from google.oauth2 import service_account
from google.cloud.aiplatform_v1beta1.types import rag as rag_types

SERVICE_ACCOUNT_PATH = "rag_agent/key.json" 

creds = service_account.Credentials.from_service_account_file(
    SERVICE_ACCOUNT_PATH, 
    scopes=['https://www.googleapis.com/auth/drive.readonly']
)
drive_service = build('drive', 'v3', credentials=creds)

# folder_names=["Appel d'offre", "Livrables"]
folder_names=["BACKMARKET"]
all_paths = []
not_found = []

# 1. Boucle pour récupérer tous les IDs
for name in folder_names:
    escaped_name = name.replace("'", "\\'")
    query = f"name = '{escaped_name}' and mimeType = 'application/vnd.google-apps.folder' and trashed = false"
    # On ajoute supportsAllDrives pour être sûr de couvrir les dossiers partagés
    response = drive_service.files().list(
        q=query, 
        fields="files(id, name)",
        supportsAllDrives=True,
        includeItemsFromAllDrives=True
    ).execute()
    
    folders = response.get('files', [])
    if folders:
        folder_id=folders[0]['id']
        path = f"https://drive.google.com/drive/u/0/folders/{folder_id}"
        all_paths.append(path)
    else:
        not_found.append(name)

# Config
chunking_config = rag.ChunkingConfig(
    chunk_size=512,        # Nombre de tokens par morceau
    chunk_overlap=50       # Chevauchement entre les morceaux
)

transformation_config = rag.TransformationConfig(
    chunking_config = rag.ChunkingConfig(
                    chunk_size=DEFAULT_CHUNK_SIZE,        # Nombre de tokens par morceau
                    chunk_overlap=DEFAULT_CHUNK_OVERLAP       # Chevauchement entre les morceaux
                )
)

# 2. Importation groupée via Vertex AI RAG        
import_result = rag.import_files(
    rag_corpus.name,
    all_paths,
    transformation_config=transformation_config,
    max_embedding_requests_per_min=DEFAULT_EMBEDDING_REQUESTS_PER_MIN,
)

files = list(rag.list_files(corpus_name=rag_corpus.name))
for f in files:
    print(f"  • {f.display_name}")




import io
import re
import PyPDF2
from docx import Document
from pptx import Presentation
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from google.oauth2 import service_account
import pandas as pd

def get_document_content(document_url: str) -> dict:
    creds = service_account.Credentials.from_service_account_file(
        SERVICE_ACCOUNT_PATH, 
        scopes=['https://www.googleapis.com/auth/drive.readonly'] # readonly suffit
    )
    
    try:
        match = re.search(r"/d/([a-zA-Z0-9-_]+)", document_url)
        if not match:
            return {"status": "error", "message": "URL non valide"}
        
        file_id = match.group(1)
        service = build('drive', 'v3', credentials=creds)

        # 1. Récupérer les métadonnées
        meta = service.files().get(
            fileId=file_id, 
            supportsAllDrives=True, 
            fields="name, mimeType"
        ).execute()
        
        mime_type = meta.get("mimeType")
        file_name = meta.get("name")
        text_content = ""

        # --- CAS 1 : DOCUMENTS NATIFS GOOGLE (Docs, Slides, Sheets) ---
        # On vérifie si c'est un format Google Apps
        if "application/vnd.google-apps" in mime_type:
            
            # Format d'export par défaut : texte brut
            export_mime = 'text/plain'
            
            # Ajustement si c'est un tableur
            if "spreadsheet" in mime_type:
                export_mime = 'text/csv'
            
            # L'export fonctionne pour Docs, Slides ET Sheets
            content = service.files().export(
                fileId=file_id, 
                mimeType=export_mime
            ).execute()
            
            text_content = content.decode('utf-8')

        # --- CAS 2 : FICHIERS BINAIRES (PDF, PPTX, WORD) ---
        else:
            request = service.files().get_media(fileId=file_id)
            file_stream = io.BytesIO()
            downloader = MediaIoBaseDownload(file_stream, request)
            done = False
            while not done:
                _, done = downloader.next_chunk()
            file_stream.seek(0)

            # Extraction PDF
            if "pdf" in mime_type:
                reader = PyPDF2.PdfReader(file_stream)
                text_content = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
            
            # Extraction PowerPoint (.pptx uniquement, les Slides sont gérés en Cas 1)
            elif "presentation" in mime_type or file_name.endswith(".pptx"):
                prs = Presentation(file_stream)
                text_runs = []
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            text_runs.append(shape.text)
                text_content = "\n".join(text_runs)
            
            # Extraction Word (.docx uniquement, les Docs sont gérés en Cas 1)
            elif "wordprocessingml" in mime_type or file_name.endswith(".docx"):
                doc = Document(file_stream)
                full_text = [para.text for para in doc.paragraphs if para.text]
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            full_text.append(cell.text)
                text_content = "\n".join(full_text)
            
            elif "spreadsheetml" in mime_type or file_name.endswith(".xlsx") or file_name.endswith(".xls"):
                # On lit le fichier Excel directement depuis le flux mémoire
                # On ne prend que la première feuille pour éviter de saturer l'agent
                df = pd.read_excel(file_stream, sheet_name=0)
                
                # Conversion en format texte (CSV ou Tabulaire)
                text_content = df.to_csv(index=False)
            
            else:
                return {"status": "error", "message": f"Format binaire non supporté : {mime_type}"}

        return {
            "status": "success",
            "title": file_name,
            "content": text_content[:10000] # Sécurité pour le contexte du modèle
        }

    except Exception as e:
        return {"status": "error", "message": str(e)}
    
document_url="https://docs.google.com/document/d/1n-x9Cu2N52_p-HFUAA48vLmsQRuEYMas/edit" #word
document_url="https://docs.google.com/presentation/d/18yzntmcbipBa-uSa7I31uzWRFkdYXl23F131vWVy6RA/edit" #gslides
document_url="https://docs.google.com/spreadsheets/d/1tinZCpubt30FVMjLJLnnKIWg9SU-n40m/edit?usp=drive_web&ouid=113892941545482869070&rtpof=true"

get_document_content(document_url)