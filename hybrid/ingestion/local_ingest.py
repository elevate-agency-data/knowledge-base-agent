"""
Local filesystem ingester — atelier model.

Ingests a local Drive-style tree into the active brand's store, projecting the
folder hierarchy onto ONE index per domain plus filterable metadata:

    <root>/<LigneProduit>/<Domaine>/[<Zone>/]<file>
            │              │         └ zone         -> meta "zone"
            │              └ domaine (leaf)          -> INDEX NAME (ascii)
            └ ligne_produit  (_Transverse -> "")     -> meta "ligne_produit"
    _meta.yaml (nearest wins, inherited up the tree) -> audience_role, matiere

Unlike the Drive path (``hybrid_add_data_auto``), this reads files straight
from disk — no service account needed — so the corpus can be tested locally
before it lands on Drive. The Drive ingester will mirror this mapping later.

Usage (index goes to the ACTIVE profile's DuckDB, so set BRAND_PROFILE):

    BRAND_PROFILE=hermes python -m scripts.ingest_local "C:/.../Rag_hermes"
"""

from __future__ import annotations

import os
from datetime import date

from hybrid.ingestion.atelier_map import (
    META_FILENAME, map_folders, as_list as _as_list,
)

_SUPPORTED_EXT = {".docx": "Docx", ".xlsx": "Xlsx", ".pdf": "PDF"}


# ── path → dimensions ─────────────────────────────────────────────────────────

def _map_path(root: str, file_path: str) -> dict | None:
    """Map a file path to (index, ligne_produit, zone). None if too shallow."""
    rel = os.path.relpath(file_path, root)
    folders = rel.split(os.sep)[:-1]          # drop the filename
    return map_folders(folders)


# ── _meta.yaml inheritance ────────────────────────────────────────────────────

def _load_meta_chain(root: str, file_path: str) -> dict:
    """Merge _meta.yaml from the file's folder up to root — nearest wins."""
    import yaml
    merged: dict = {}
    cur = os.path.dirname(file_path)
    root = os.path.abspath(root)
    chain = []
    while True:
        chain.append(cur)
        if os.path.abspath(cur) == root:
            break
        parent = os.path.dirname(cur)
        if parent == cur:
            break
        cur = parent
    # walk from root downward so nearer folders overwrite farther ones
    for folder in reversed(chain):
        mp = os.path.join(folder, META_FILENAME)
        if os.path.isfile(mp):
            try:
                data = yaml.safe_load(open(mp, encoding="utf-8")) or {}
                if isinstance(data, dict):
                    merged.update(data)
            except Exception:
                pass
    return merged


# ── text extraction ───────────────────────────────────────────────────────────

def _extract(file_path: str, ext: str) -> str:
    if ext == ".docx":
        from docx import Document
        doc = Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    if ext == ".pdf":
        try:
            import fitz
        except Exception:
            import pymupdf as fitz
        doc = fitz.open(file_path)
        return "\n".join(page.get_text() for page in doc)
    return ""


# ── ingestion ─────────────────────────────────────────────────────────────────

def ingest_local_tree(root: str, chunk_strategy: str = "fixed") -> dict:
    """Ingest every supported file under *root* into the active store.

    Returns a summary dict. Idempotent-ish: re-running INSERT OR REPLACEs by id,
    but ids are content-derived per run — call on a fresh index for clean state.
    """
    from hybrid.stores import get_store
    from hybrid.embeddings import get_embedding_model
    from hybrid.config import (
        DEFAULT_EMBEDDING_MODEL, CHUNK_SIZE, CHUNK_OVERLAP,
    )
    from hybrid.ingestion.chunker import chunk_text
    from hybrid.ingestion.metadata import build_metadata, detect_language

    root = os.path.abspath(root)
    if not os.path.isdir(root):
        return {"status": "error", "message": f"Not a directory: {root}"}

    store = get_store()
    embedder = get_embedding_model(DEFAULT_EMBEDDING_MODEL)
    emb_dim = embedder.get_dimension()
    chunk_params = ({"size": CHUNK_SIZE, "overlap": CHUNK_OVERLAP}
                    if chunk_strategy == "fixed" else {})

    per_index: dict[str, int] = {}
    files_done = 0
    files_skipped: list[str] = []
    seen_indexes: set[str] = set()

    for dirpath, _dirs, files in os.walk(root):
        for fname in sorted(files):
            ext = os.path.splitext(fname)[1].lower()
            if ext not in _SUPPORTED_EXT:
                continue
            fpath = os.path.join(dirpath, fname)
            mp = _map_path(root, fpath)
            if not mp:
                files_skipped.append(f"{fname} (too shallow)")
                continue

            index_name = mp["index"]
            if index_name not in seen_indexes:
                store.initialize(index_name,
                                 embedding_model=DEFAULT_EMBEDDING_MODEL,
                                 chunk_strategy=chunk_strategy)
                seen_indexes.add(index_name)

            text = _extract(fpath, ext)
            if not text.strip():
                files_skipped.append(f"{fname} (empty)")
                continue

            chunks = chunk_text(text, strategy=chunk_strategy, **chunk_params)
            if not chunks:
                files_skipped.append(f"{fname} (no chunks)")
                continue

            meta_yaml = _load_meta_chain(root, fpath)
            audience_role = _as_list(meta_yaml.get("audience_role"))
            matiere = _as_list(meta_yaml.get("matiere"))
            langue = detect_language(text[:5000])
            try:
                mtime = date.fromtimestamp(os.path.getmtime(fpath)).isoformat()
            except Exception:
                mtime = ""

            file_info = {
                "file_name": fname,
                "file_type": _SUPPORTED_EXT[ext],
                "source_url": fpath,
                "author": "",
                "created_at": mtime,
                "updated_at": mtime,
            }

            texts = [c["content"] for c in chunks]
            embeddings = embedder.embed_documents(texts)

            records = []
            for chunk, emb in zip(chunks, embeddings):
                rec = build_metadata(
                    file_info=file_info,
                    chunk=chunk,
                    embedding_model=DEFAULT_EMBEDDING_MODEL,
                    index_name=index_name,
                    embedding_dim=emb_dim,
                    doc_language=langue,
                    doc_domaine=mp["domaine"].upper(),
                    ligne_produit=mp["ligne_produit"],
                    zone=mp["zone"],
                    audience_role=audience_role,
                    matiere=matiere,
                )
                rec["embedding"] = emb
                records.append(rec)

            store.insert_chunks(records)
            per_index[index_name] = per_index.get(index_name, 0) + len(records)
            files_done += 1
            print(f"  [{index_name:16s}] {fname}  "
                  f"(ligne={mp['ligne_produit'] or '-'} zone={mp['zone'] or '-'} "
                  f"role={audience_role or '-'} +{len(records)} chunks)")

    return {
        "status": "success",
        "root": root,
        "indexes": sorted(per_index),
        "chunks_per_index": per_index,
        "files_ingested": files_done,
        "files_skipped": files_skipped,
        "total_chunks": sum(per_index.values()),
    }
