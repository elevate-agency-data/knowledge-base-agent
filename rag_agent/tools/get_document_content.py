import io
import re
import PyPDF2
from docx import Document
from pptx import Presentation
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from google.oauth2 import service_account
import pandas as pd
from ..config import (
    SERVICE_ACCOUNT_PATH 
)

def get_document_content(document_url: str) -> dict:
    """
    Retrieves the text content of a Google Drive document given its full URL.
    Args:
        document_url: The full URL of the document (e.g., 'https://docs.google.com/...')
    Returns:
        A dictionary containing the status, title, and text content.
    """
    creds = service_account.Credentials.from_service_account_file(
        SERVICE_ACCOUNT_PATH, 
        scopes=['https://www.googleapis.com/auth/drive.readonly'] # readonly suffit
    )
    
    try:
        match = re.search(r"/d/([a-zA-Z0-9-_]+)", document_url)
        if not match:
            return {"status": "error", "message": "URL non valide"}
        
        file_id = match.group(1)
        service = build('drive', 'v3', credentials=creds)

        # 1. Récupérer les métadonnées
        meta = service.files().get(
            fileId=file_id, 
            supportsAllDrives=True, 
            fields="name, mimeType"
        ).execute()
        
        mime_type = meta.get("mimeType")
        file_name = meta.get("name")
        text_content = ""

        # --- CAS 1 : DOCUMENTS NATIFS GOOGLE (Docs, Slides, Sheets) ---
        # On vérifie si c'est un format Google Apps
        if "application/vnd.google-apps" in mime_type:
            
            # Format d'export par défaut : texte brut
            export_mime = 'text/plain'
            
            # Ajustement si c'est un tableur
            if "spreadsheet" in mime_type:
                export_mime = 'text/csv'
            
            # L'export fonctionne pour Docs, Slides ET Sheets
            content = service.files().export(
                fileId=file_id, 
                mimeType=export_mime
            ).execute()
            
            text_content = content.decode('utf-8')

        # --- CAS 2 : FICHIERS BINAIRES (PDF, PPTX, WORD) ---
        else:
            request = service.files().get_media(fileId=file_id)
            file_stream = io.BytesIO()
            downloader = MediaIoBaseDownload(file_stream, request)
            done = False
            while not done:
                _, done = downloader.next_chunk()
            file_stream.seek(0)

            # Extraction PDF
            if "pdf" in mime_type:
                reader = PyPDF2.PdfReader(file_stream)
                text_content = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
            
            # Extraction PowerPoint (.pptx uniquement, les Slides sont gérés en Cas 1)
            elif "presentation" in mime_type or file_name.endswith(".pptx"):
                prs = Presentation(file_stream)
                text_runs = []
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            text_runs.append(shape.text)
                text_content = "\n".join(text_runs)
            
            # Extraction Word 
            elif "wordprocessingml" in mime_type or file_name.endswith(".docx"):
                doc = Document(file_stream)
                full_text = [para.text for para in doc.paragraphs if para.text]
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            full_text.append(cell.text)
                text_content = "\n".join(full_text)
            
            elif "spreadsheetml" in mime_type or file_name.endswith(".xlsx") or file_name.endswith(".xls"):
                # On lit le fichier Excel directement depuis le flux mémoire
                # On ne prend que la première feuille pour éviter de saturer l'agent
                df = pd.read_excel(file_stream, sheet_name=0)
                
                text_content = df.to_csv(index=False)
            
            else:
                return {"status": "error", "message": f"Format binaire non supporté : {mime_type}"}

        return {
            "status": "success",
            "title": file_name,
            "content": text_content[:10000] # Sécurité pour le contexte du modèle
        }

    except Exception as e:
        return {"status": "error", "message": str(e)}